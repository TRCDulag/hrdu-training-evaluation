# Writes the empty Sheet template Register/HRDU Training Evaluation Register.xlsx from the tab headers in Code.gs.
import json, pathlib, subprocess
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
root = pathlib.Path(__file__).resolve().parent.parent
js = "const vm=require('vm');const {loadBackend}=require('./dev/mock-gas');const b=loadBackend();const g=k=>vm.runInContext(k,b.ctx);" \
     "console.log(JSON.stringify({T:g('TRAINING_HEADERS'),R:g('RESP_HEADERS'),P:g('RP_HEADERS'),U:g('USER_HEADERS'),A:g('AUDIT_HEADERS'),V:g('VERSION')}))"
h = json.loads(subprocess.check_output(["node", "-e", js], cwd=root))
wb = Workbook(); wb.remove(wb.active)
for name, cols in [("Trainings", h["T"]), ("Responses", h["R"]), ("Resource Person Ratings", h["P"]), ("Users", h["U"]), ("Audit Trail", h["A"])]:
    ws = wb.create_sheet(name)
    ws.append(cols)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF"); c.fill = PatternFill("solid", fgColor="1F2A45")
    for i, col in enumerate(cols, 1):
        ws.column_dimensions[ws.cell(1, i).column_letter].width = min(40, max(12, len(col) + 2))
    ws.freeze_panes = "A2"
ws = wb.create_sheet("About")
for line in ["HRDU Training Evaluation Register (Training Evaluation Tool, Level 1)", "Version " + h["V"],
             "Upload to Google Drive, open, File > Save as Google Sheets, then follow the README (Extensions > Apps Script).",
             "Do not type in the tabs by hand. The app writes to them."]:
    ws.append([line])
out = root / "Register" / "HRDU Training Evaluation Register.xlsx"
wb.save(out); print(out)
